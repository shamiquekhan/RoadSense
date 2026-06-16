#!/usr/bin/env python3
"""
scripts/generate_findings.py
Generate a 5-page Findings Summary Word document from a template.

Usage:
  python scripts/generate_findings.py --output outputs/reports/RoadSense_Findings_Summary.docx

The script will try to read processed outputs (GeoPackage or GeoJSON) to
replace placeholders. If `python-docx` is not installed, it will write a
plain-text fallback at the same path with .txt extension.
"""
import argparse
import json
from pathlib import Path
import sys

TEMPLATE_PATH = Path("outputs/reports/RoadSense_Findings_Summary_template.txt")


def read_template():
    if not TEMPLATE_PATH.exists():
        print(f"Template not found: {TEMPLATE_PATH}")
        sys.exit(1)
    return TEMPLATE_PATH.read_text(encoding="utf-8")


def gather_metrics():
    # Try common output locations to compute simple metrics
    metrics = {
        "TOTAL_SEGMENTS": "TBD",
        "N_CRITICAL": "TBD",
        "PCT_CRITICAL": "TBD",
        "N_HIGH": "TBD",
        "PCT_HIGH": "TBD",
        "N_MEDIUM": "TBD",
        "PCT_MEDIUM": "TBD",
        "N_LOW": "TBD",
        "PCT_LOW": "TBD",
        "PCT_V85_ABOVE_LIMIT": "TBD",
        "MEDIAN_SPEED_GAP": "TBD",
        "PCT_NO_FOOTPATH": "TBD",
    }
    try:
        import geopandas as gpd
        gpkg = Path("outputs/all_segments_scored.gpkg")
        geojson = Path("outputs/geojson/all_segments_scored.geojson")
        if gpkg.exists():
            gdf = gpd.read_file(gpkg, layer="segments")
        elif geojson.exists():
            gdf = gpd.read_file(geojson)
        else:
            # try processed master
            master = Path("data/processed/segments_master.gpkg")
            if master.exists():
                gdf = gpd.read_file(master, layer="segments")
            else:
                return metrics

        metrics["TOTAL_SEGMENTS"] = int(len(gdf))
        # risk tier column assumed `risk_tier` with values Critical/High/Medium/Low
        if "risk_tier" in gdf.columns:
            for tier in ["Critical", "High", "Medium", "Low"]:
                n = int((gdf.get("risk_tier", "").astype(str) == tier).sum())
                pct = round(100 * n / len(gdf), 1) if len(gdf) else 0.0
                key_n = f"N_{tier.upper()}" if tier != "Critical" else "N_CRITICAL"
                key_pct = f"PCT_{tier.upper()}" if tier != "Critical" else "PCT_CRITICAL"
                metrics[key_n] = n
                metrics[key_pct] = f"{pct}%"

        # simple speed metrics
        if "v85" in gdf.columns and "posted_limit" in gdf.columns:
            above = gdf[~gdf["v85"].isna() & ~gdf["posted_limit"].isna() & (gdf["v85"] > gdf["posted_limit"] + 10)]
            metrics["PCT_V85_ABOVE_LIMIT"] = f"{round(100 * len(above) / len(gdf),1)}%" if len(gdf) else "0%"
            metrics["MEDIAN_SPEED_GAP"] = round((gdf["v85"] - gdf["posted_limit"]).median(skipna=True),1)

        # imagery protective feature proxy
        if "footpath_detected" in gdf.columns:
            no_fp = int((gdf["footpath_detected"].fillna(0) == 0).sum())
            metrics["PCT_NO_FOOTPATH"] = f"{round(100 * no_fp / len(gdf),1)}%"

    except Exception:
        # geopandas or data not available — leave TBD
        pass

    return metrics


def fill_template(template: str, metrics: dict) -> str:
    out = template
    for k, v in metrics.items():
        out = out.replace(f"[{k}]", str(v))
    return out


def write_docx(text: str, out_path: Path):
    try:
        from docx import Document
        doc = Document()
        # split into pages by double newline separation marker 'Page '
        pages = text.split('\n\nPage') if '\n\nPage' in text else [text]
        for p in pages:
            # ensure no extremely long lines break
            for line in p.strip().split('\n'):
                doc.add_paragraph(line)
            doc.add_page_break()
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out_path)
        print(f"Wrote Word doc: {out_path}")
        return True
    except Exception as e:
        print("python-docx not available or error creating .docx:", e)
        return False


def write_text_fallback(text: str, out_path: Path):
    out_path_txt = out_path.with_suffix(".txt")
    out_path_txt.parent.mkdir(parents=True, exist_ok=True)
    out_path_txt.write_text(text, encoding="utf-8")
    print(f"Wrote text fallback: {out_path_txt}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="outputs/reports/RoadSense_Findings_Summary.docx")
    args = parser.parse_args()
    out_path = Path(args.output)

    template = read_template()
    metrics = gather_metrics()
    filled = fill_template(template, metrics)

    ok = write_docx(filled, out_path)
    if not ok:
        write_text_fallback(filled, out_path)


if __name__ == "__main__":
    main()
